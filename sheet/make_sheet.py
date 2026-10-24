"""Generate Work_Log_Backup.xlsx: a monthly work log for Google Sheets / Excel.

One tab per month (date, day, three In/Out pairs in 12-hour time, auto working
hours, remarks) plus a Summary tab. Run:  python sheet/make_sheet.py
"""

import calendar
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as col
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent / "Work_Log_Backup.xlsx"
MONTHS = [(2026, m) for m in range(10, 13)] + [(2027, m) for m in range(1, 10)]  # Oct 2026 - Sep 2027
PAIRS = 3  # In/Out time pairs per day

ACCENT, HEAD_BG, CALC_BG, WEEKEND = "4F46E5", "EEF2FF", "F3F4F6", "FEF3C7"
thin = Side(style="thin", color="D1D5DB")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
FIRST = 5  # first data row

# Column layout: Date, Day, In1, Out1, ..., InN, OutN, Working Hours, Remarks
TIME_COLS = list(range(3, 3 + 2 * PAIRS))
HOURS_COL = 3 + 2 * PAIRS
REMARKS_COL = HOURS_COL + 1
HEADERS = ["Date", "Day"]
for i in range(1, PAIRS + 1):
    HEADERS += [f"In Time ({i})", f"Out Time ({i})"]
HEADERS += ["Working Hours", "Remarks (what I worked on)"]
WIDTHS = [14, 12] + [13] * (2 * PAIRS) + [15, 60]
LAST_COL = col(REMARKS_COL)


def header_cell(cell, text):
    cell.value = text
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=ACCENT)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BOX


def total_cell(cell, value, fmt):
    cell.value = value
    cell.number_format = fmt
    cell.font = Font(bold=True, color=ACCENT)
    cell.fill = PatternFill("solid", fgColor=HEAD_BG)
    cell.border = BOX
    cell.alignment = Alignment(horizontal="center")


def hours_formula(r):
    # Sum each complete In/Out pair; MOD(...,1) handles a pair that crosses midnight.
    first, last = col(TIME_COLS[0]), col(TIME_COLS[-1])
    parts = []
    for i in range(PAIRS):
        a, b = col(TIME_COLS[2 * i]), col(TIME_COLS[2 * i + 1])
        parts.append(f'IF(AND({a}{r}<>"",{b}{r}<>""),MOD({b}{r}-{a}{r},1),0)')
    return f'=IF(COUNT({first}{r}:{last}{r})=0,"",{"+".join(parts)})'


def month_sheet(wb, year, month):
    ws = wb.create_sheet(date(year, month, 1).strftime("%b %Y"))
    days = calendar.monthrange(year, month)[1]
    last = FIRST + days - 1

    ws.merge_cells(f"A1:{LAST_COL}1")
    ws["A1"] = f"Work Log — {date(year, month, 1):%B %Y}"
    ws["A1"].font = Font(size=16, bold=True, color="111827")
    ws.merge_cells(f"A2:{LAST_COL}2")
    ws["A2"] = (
        "Type times in 12-hour format with AM/PM, e.g. 9:00 AM or 6:30 PM (6:30 without PM is read as morning). "
        "Working Hours adds up every In/Out pair automatically; leave unused pairs empty."
    )
    ws["A2"].font = Font(size=9, italic=True, color="6B7280")

    for c, (h, w) in enumerate(zip(HEADERS, WIDTHS), start=1):
        header_cell(ws.cell(4, c), h)
        ws.column_dimensions[col(c)].width = w
    ws.row_dimensions[4].height = 30

    for i in range(days):
        r = FIRST + i
        ws.cell(r, 1, date(year, month, i + 1)).number_format = "dd mmm yyyy"
        ws.cell(r, 2, f'=TEXT(A{r},"dddd")')
        for c in TIME_COLS:
            ws.cell(r, c).number_format = "h:mm AM/PM"
        ws.cell(r, HOURS_COL, hours_formula(r)).number_format = "[h]:mm"
        for c in range(1, REMARKS_COL + 1):
            cell = ws.cell(r, c)
            cell.border = BOX
            cell.alignment = Alignment(
                horizontal="left" if c == REMARKS_COL else "center", vertical="center", wrap_text=c == REMARKS_COL
            )
            if c in (1, 2, HOURS_COL):
                cell.fill = PatternFill("solid", fgColor=CALC_BG)
        ws.cell(r, HOURS_COL).font = Font(bold=True)

    # Friday & Saturday shaded (editable in Format > Conditional formatting).
    ws.conditional_formatting.add(
        f"A{FIRST}:{LAST_COL}{last}",
        FormulaRule(
            formula=[f"OR(WEEKDAY($A{FIRST})=6,WEEKDAY($A{FIRST})=7)"],
            fill=PatternFill("solid", fgColor=WEEKEND, bgColor=WEEKEND),
        ),
    )

    dv = DataValidation(
        type="time", operator="between", formula1="0", formula2="0.999988426", allow_blank=True,
        showErrorMessage=True, errorTitle="Not a time", error="Enter a time with AM/PM, e.g. 9:00 AM or 6:30 PM",
    )
    dv.add(f"{col(TIME_COLS[0])}{FIRST}:{col(TIME_COLS[-1])}{last}")
    ws.add_data_validation(dv)

    h = col(HOURS_COL)
    t = last + 2
    totals = [
        ("Total working hours", f"=SUM({h}{FIRST}:{h}{last})", "[h]:mm"),
        ("Total hours (decimal)", f"=ROUND({h}{t}*24,2)", "0.00"),
        ("Days worked", f'=COUNTIF({h}{FIRST}:{h}{last},">0")', "0"),
        ("Average per day worked", f'=IF({h}{t + 2}=0,"",{h}{t}/{h}{t + 2})', "[h]:mm"),
    ]
    label_from = HOURS_COL - 3
    for k, (label, formula, fmt) in enumerate(totals):
        r = t + k
        ws.merge_cells(start_row=r, start_column=label_from, end_row=r, end_column=HOURS_COL - 1)
        lc = ws.cell(r, label_from, label)
        lc.font = Font(bold=True)
        lc.alignment = Alignment(horizontal="right")
        total_cell(ws.cell(r, HOURS_COL), formula, fmt)

    ws.freeze_panes = f"C{FIRST}"
    ws.print_title_rows = "4:4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws.title, t


def summary_sheet(wb, months):
    s = wb.create_sheet("Summary", 0)
    s.merge_cells("A1:D1")
    s["A1"] = "Work Log — Summary"
    s["A1"].font = Font(size=16, bold=True, color="111827")
    for c, (head, w) in enumerate(
        zip(["Month", "Working Hours", "Hours (decimal)", "Days Worked"], [16, 16, 16, 14]), start=1
    ):
        header_cell(s.cell(3, c), head)
        s.column_dimensions[col(c)].width = w

    h = col(HOURS_COL)
    for i, (name, t) in enumerate(months):
        r = 4 + i
        ref = f"'{name}'!"
        cells = [(name, None), (f"={ref}{h}{t}", "[h]:mm"), (f"={ref}{h}{t + 1}", "0.00"), (f"={ref}{h}{t + 2}", "0")]
        for c, (v, fmt) in enumerate(cells, start=1):
            cell = s.cell(r, c, v)
            cell.border = BOX
            cell.alignment = Alignment(horizontal="center")
            if fmt:
                cell.number_format = fmt

    end = 3 + len(months)
    r = end + 1
    s.cell(r, 1, "Total").font = Font(bold=True)
    s.cell(r, 1).border = BOX
    for c, fmt in [(2, "[h]:mm"), (3, "0.00"), (4, "0")]:
        total_cell(s.cell(r, c), f"=SUM({col(c)}4:{col(c)}{end})", fmt)
    s.freeze_panes = "A4"


def main():
    wb = Workbook()
    wb.remove(wb.active)
    months = [month_sheet(wb, y, m) for y, m in MONTHS]
    summary_sheet(wb, months)

    # Open on the current month if it's in range.
    current = date.today().strftime("%b %Y")
    wb.active = wb.sheetnames.index(current) if current in wb.sheetnames else 0
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws is wb.active
    wb.save(OUT)
    print(f"Saved {OUT}")


if __name__ == "__main__":
    main()
